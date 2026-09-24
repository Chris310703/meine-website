"""Hochgeladene Lernunterlagen speichern und Text daraus lesen (PDF, PowerPoint, Text)."""

from __future__ import annotations

import re
import uuid
from pathlib import Path

from .. import config

SUPPORTED = {".pdf", ".pptx", ".txt", ".md"}


class DocumentError(Exception):
    pass


def safe_name(filename: str) -> str:
    name = Path(filename or "datei").name
    name = re.sub(r"[^\w.\-äöüÄÖÜß ]", "_", name).strip() or "datei"
    return name[:120]


def subject_dir(subject_id: int) -> Path:
    path = config.UPLOAD_DIR / f"fach_{subject_id}"
    path.mkdir(parents=True, exist_ok=True)
    return path


def store(subject_id: int, filename: str, data: bytes) -> Path:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        hint = " Alte .ppt-Dateien bitte als PDF exportieren." if suffix == ".ppt" else ""
        raise DocumentError(f"Dateityp {suffix or '(ohne Endung)'} wird nicht unterstützt. Erlaubt: PDF, PPTX, TXT, MD.{hint}")
    if len(data) > 60 * 1024 * 1024:
        raise DocumentError("Die Datei ist größer als 60 MB.")
    path = subject_dir(subject_id) / f"{uuid.uuid4().hex[:8]}_{safe_name(filename)}"
    path.write_bytes(data)
    return path


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            parts = []
            for i, page in enumerate(reader.pages, start=1):
                text = (page.extract_text() or "").strip()
                if text:
                    parts.append(f"[Seite {i}]\n{text}")
            return "\n\n".join(parts)
        if suffix == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = []
            for i, slide in enumerate(prs.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                        t = shape.text_frame.text.strip()
                        if t:
                            texts.append(t)
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        texts.append(f"Notizen: {notes}")
                if texts:
                    parts.append(f"[Folie {i}]\n" + "\n".join(texts))
            return "\n\n".join(parts)
        return path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # beschädigte Datei o. Ä.
        raise DocumentError(f"Text konnte nicht gelesen werden: {exc}") from exc


def text_cache_path(path: Path) -> Path:
    return path.with_suffix(path.suffix + ".txt")


def cached_text(path: Path) -> str:
    cache = text_cache_path(path)
    if cache.exists():
        return cache.read_text(encoding="utf-8", errors="replace")
    text = extract_text(path)
    cache.write_text(text, encoding="utf-8")
    return text


def delete_files(path: Path) -> None:
    for p in (path, text_cache_path(path)):
        try:
            p.unlink()
        except FileNotFoundError:
            pass
